"""SÜDEV Fruit Ninja - Python sunucusu (arkaplan-excel.js'in karşılığı)."""
import json
import os
from datetime import datetime
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer

from openpyxl import Workbook, load_workbook

HOST = '127.0.0.1'
PORT = 3000
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, 'puanlar.xlsx')
JSON_BACKUP_PATH = os.path.join(BASE_DIR, 'puanlar.json')
HEADERS = ['İsim', 'Puan', 'Tarih', 'Saat']


def to_number(v):
    try:
        return float(v) if '.' in str(v) else int(v)
    except (TypeError, ValueError):
        return 0


def normalize_entry(r):
    raw_name = r.get('isim') or r.get('instagram') or r.get('İsim') or r.get('Instagram') or ''
    raw_name = str(raw_name)
    isim = raw_name if raw_name.startswith('@') else '@' + raw_name
    puan = to_number(next((r[k] for k in ('puan', 'score', 'Puan') if r.get(k) is not None), 0))

    tarih = r.get('tarih') or r.get('Tarih') or ''
    saat = r.get('saat') or r.get('Saat') or ''

    combined = r.get('dateFormatted') or r.get('Tarih / Saat')
    if (not tarih or not saat) and combined:
        parts = str(combined).split()
        tarih = tarih or (parts[0] if parts else '')
        saat = saat or (parts[1] if len(parts) > 1 else '')

    now = datetime.now()
    tarih = tarih or now.strftime('%d.%m.%Y')
    saat = saat or now.strftime('%H:%M:%S')

    return {'isim': isim, 'puan': puan, 'tarih': str(tarih), 'saat': str(saat)}


def sort_rows_desc(rows):
    rows.sort(key=lambda r: to_number(r.get('puan')), reverse=True)
    return rows


def read_rows():
    try:
        if os.path.exists(JSON_BACKUP_PATH):
            with open(JSON_BACKUP_PATH, encoding='utf-8') as f:
                data = json.load(f)
            if isinstance(data, list):
                return sort_rows_desc([
                    normalize_entry(r) for r in data
                    if isinstance(r, dict) and (r.get('isim') or r.get('instagram'))
                ])
    except Exception:
        pass

    try:
        if os.path.exists(EXCEL_PATH):
            ws = load_workbook(EXCEL_PATH, read_only=True).worksheets[0]
            rows = list(ws.iter_rows(values_only=True))
            if rows:
                header = [str(h) if h is not None else '' for h in rows[0]]
                raw = [dict(zip(header, row)) for row in rows[1:]]
                return sort_rows_desc([
                    normalize_entry(r) for r in raw
                    if (r.get('İsim') and r.get('İsim') != '-')
                    or (r.get('Instagram') and r.get('Instagram') != '-')
                ])
    except Exception:
        pass

    return []


def write_excel(rows):
    clean = sort_rows_desc([normalize_entry(r) for r in rows])
    with open(JSON_BACKUP_PATH, 'w', encoding='utf-8') as f:
        json.dump(clean, f, ensure_ascii=False, indent=2)

    # Excel'de puana göre çoktan aza sıralı liste (Sadece İsim, Puan, Tarih, Saat)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Puanlar'
    ws.append(HEADERS)
    for item in clean:
        ws.append([item['isim'], item['puan'], item['tarih'], item['saat']])
    for col, width in zip('ABCD', (28, 12, 16, 14)):
        ws.column_dimensions[col].width = width

    try:
        wb.save(EXCEL_PATH)
        return True
    except OSError as err:
        print('puanlar.xlsx şu an Excel programında açık olabilir:', err)
        return False


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=BASE_DIR, **kwargs)

    def end_headers(self):
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, DELETE, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        super().end_headers()

    def send_json(self, payload, status=200):
        body = json.dumps(payload, ensure_ascii=False).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def route(self):
        return self.path.split('?', 1)[0]

    def do_OPTIONS(self):
        self.send_response(200)
        self.end_headers()

    def do_GET(self):
        if self.route() == '/api/ping':
            return self.send_json({'ok': True})
        if self.route() == '/api/scores':
            return self.send_json({'ok': True, 'scores': read_rows(), 'excelPath': EXCEL_PATH})
        return super().do_GET()

    def do_POST(self):
        if self.route() != '/api/scores':
            return self.send_json({'ok': False, 'error': 'Bulunamadı'}, 404)
        length = int(self.headers.get('Content-Length') or 0)
        try:
            body = json.loads(self.rfile.read(length) or b'{}')
        except ValueError:
            body = {}
        if not isinstance(body, dict) or not (body.get('isim') or body.get('instagram')):
            return self.send_json({'ok': False, 'error': 'İsim gerekli'}, 400)
        rows = read_rows()
        rows.append(normalize_entry(body))
        sort_rows_desc(rows)
        saved = write_excel(rows)
        self.send_json({'ok': True, 'saved': saved, 'total': len(rows), 'scores': rows})

    def do_DELETE(self):
        if self.route() != '/api/scores':
            return self.send_json({'ok': False, 'error': 'Bulunamadı'}, 404)
        write_excel([])
        self.send_json({'ok': True, 'scores': []})


if __name__ == '__main__':
    if not os.path.exists(EXCEL_PATH):
        write_excel(read_rows())
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f'SÜDEV Fruit Ninja Uygulama Servisi Aktif: {EXCEL_PATH}')
    print(f'http://{HOST}:{PORT}/')
    server.serve_forever()
